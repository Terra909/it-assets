from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib.auth import login, logout
from django.contrib.auth.forms import AuthenticationForm, UserCreationForm
from django.contrib.auth.models import Group, User
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponseForbidden, HttpResponse
from django.db import models
from django.db.models import Q
from django.utils import timezone
from decimal import Decimal, InvalidOperation
from datetime import datetime
import qrcode
from io import BytesIO
import base64

from .access import get_user_access
from .models import Asset, AssetType, Status, Location, RepairRequest, ActionLog, BYODResource, InventorySession, InventoryItem


PRIVILEGED_GROUP_PERMISSION = 'accounts.manage_privileged_employees'
PRIVILEGED_VIEW_PERMISSION = 'accounts.view_privileged_employees'


def get_manageable_groups(profile):
    groups = Group.objects.order_by('name')
    if profile.can_manage_privileged_employees():
        return groups
    return groups.exclude(permissions__codename='manage_privileged_employees').distinct()


def get_assignable_groups(profile):
    groups = get_manageable_groups(profile)
    if profile.is_manager() and not profile.can_manage_privileged_employees():
        return groups.filter(name='Сотрудник')
    return groups


def is_privileged_user(target_user):
    return (
        target_user.is_superuser
        or target_user.has_perm(PRIVILEGED_GROUP_PERMISSION)
        or target_user.has_perm(PRIVILEGED_VIEW_PERMISSION)
    )


def is_employee_or_unassigned_user(target_user):
    primary_group = get_primary_group(target_user)
    return primary_group is None or primary_group.name == 'Сотрудник'


def get_visible_employees_queryset(profile):
    users = User.objects.prefetch_related('groups').all()
    if profile.can_view_privileged_employees():
        return users
    return users.exclude(groups__permissions__codename='manage_privileged_employees').exclude(is_superuser=True).distinct()


def get_primary_group(user):
    return user.groups.order_by('name').first()


def get_primary_group_name(user):
    primary_group = get_primary_group(user)
    return primary_group.name if primary_group else ''


def log_action(user, action_type, description='', object_id=None):
    """Функция для записи действия в журнал"""
    ActionLog.objects.create(
        user=user,
        action_type=action_type,
        description=description,
        object_id=object_id
    )


def deny_access(message="Недостаточно прав"):
    return HttpResponseForbidden(message)


def require_permission(profile, permission_checker, message="Недостаточно прав для выполнения действия"):
    checker = getattr(profile, permission_checker)
    if not checker():
        return deny_access(message)
    return None


def can_transfer_asset(profile, asset):
    return profile.can_manage_assets() or asset.employee_id == profile.user.id


def build_inventory_items_for_session(session):
    assets = Asset.objects.select_related('employee', 'location', 'status').all()
    if session.location_id:
        assets = assets.filter(location_id=session.location_id)

    items = [
        InventoryItem(
            session=session,
            asset=asset,
            expected_employee=asset.employee,
            expected_location=asset.location,
            expected_status=asset.status,
        )
        for asset in assets
    ]
    InventoryItem.objects.bulk_create(items)


def get_inventory_report_sessions_queryset(session_search=''):
    sessions = InventorySession.objects.select_related('location', 'created_by').all()
    if session_search:
        sessions = sessions.filter(
            Q(title__icontains=session_search)
            | Q(notes__icontains=session_search)
            | Q(location__name__icontains=session_search)
            | Q(created_by__username__icontains=session_search)
            | Q(created_by__first_name__icontains=session_search)
            | Q(created_by__last_name__icontains=session_search)
        )
    return sessions.order_by('-created_at')


def get_inventory_report_items_queryset(session, query='', result_filter=''):
    items = session.items.select_related(
        'asset',
        'expected_employee', 'expected_location', 'expected_status',
        'actual_employee', 'actual_location', 'actual_status',
        'checked_by'
    ).all()

    if query:
        items = items.filter(
            Q(asset__asset_name__icontains=query)
            | Q(asset__inventory_number__icontains=query)
            | Q(asset__serial_number__icontains=query)
            | Q(comment__icontains=query)
            | Q(expected_employee__username__icontains=query)
            | Q(expected_employee__first_name__icontains=query)
            | Q(expected_employee__last_name__icontains=query)
            | Q(actual_employee__username__icontains=query)
            | Q(actual_employee__first_name__icontains=query)
            | Q(actual_employee__last_name__icontains=query)
            | Q(expected_location__name__icontains=query)
            | Q(actual_location__name__icontains=query)
            | Q(expected_status__name__icontains=query)
            | Q(actual_status__name__icontains=query)
        )

    valid_results = {choice[0] for choice in InventoryItem.ResultChoices.choices}
    if result_filter in valid_results:
        items = items.filter(result=result_filter)

    return items.order_by('asset__inventory_number', 'asset__asset_name')


def build_inventory_report_context(selected_session, items):
    if not selected_session:
        return {
            'items_count': 0,
            'checked_count': 0,
            'pending_count': 0,
            'matched_count': 0,
            'mismatch_count': 0,
            'missing_count': 0,
        }

    total = selected_session.items.count()
    checked = selected_session.items.exclude(result=InventoryItem.ResultChoices.PENDING).count()
    return {
        'items_count': items.count(),
        'checked_count': checked,
        'pending_count': max(total - checked, 0),
        'matched_count': selected_session.items.filter(result=InventoryItem.ResultChoices.MATCHED).count(),
        'mismatch_count': selected_session.items.filter(result=InventoryItem.ResultChoices.MISMATCH).count(),
        'missing_count': selected_session.items.filter(result=InventoryItem.ResultChoices.MISSING).count(),
    }


def normalize_header(value):
    if value is None:
        return ''
    return str(value).strip().lower().replace('ё', 'е')


def normalize_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def parse_excel_date(value):
    if value in (None, ''):
        return None
    if hasattr(value, 'date') and not isinstance(value, str):
        try:
            return value.date()
        except Exception:
            pass
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return None
        for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y'):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    raise ValueError('Некорректная дата покупки')


def parse_excel_price(value):
    if value in (None, ''):
        return None
    if isinstance(value, str):
        value = value.strip().replace(' ', '').replace(',', '.')
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError('Некорректная цена')


def build_asset_import_workbook():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Assets'

    headers = [
        'asset_name',
        'inventory_number',
        'serial_number',
        'asset_type',
        'status',
        'employee_username',
        'location',
        'purchase_date',
        'price',
        'description',
    ]
    worksheet.append(headers)

    sample_rows = [
        ['Lenovo ThinkPad T14 Gen 3', 'NB-2026-001', 'SN-LEN-001', 'Ноутбук', 'В эксплуатации', '', 'Главный офис', '2026-01-15', '125000.00', 'Тестовый ноутбук для проверки импорта'],
        ['Dell P2723DE', 'MN-2026-002', 'SN-DEL-002', 'Монитор', 'На складе', '', 'Склад 1', '2026-02-01', '28990.50', 'Пример строки с монитором'],
    ]

    for row in sample_rows:
        worksheet.append(row)

    header_fill = PatternFill(fill_type='solid', fgColor='1F2937')
    header_font = Font(color='FFFFFF', bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border

    widths = {'A': 28, 'B': 22, 'C': 22, 'D': 18, 'E': 18, 'F': 22, 'G': 20, 'H': 16, 'I': 14, 'J': 40}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    return workbook


def import_assets_from_workbook(workbook):
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError('Файл Excel пустой')

    headers = [normalize_header(value) for value in rows[0]]
    required_headers = {'inventory_number', 'asset_type', 'status'}
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise ValueError('В Excel отсутствуют обязательные колонки: ' + ', '.join(missing_headers))

    header_index = {header: index for index, header in enumerate(headers) if header}
    created_count = 0
    updated_count = 0
    skipped_rows = []

    for row_number, row in enumerate(rows[1:], start=2):
        if not any(cell not in (None, '') for cell in row):
            continue

        def get_value(column_name):
            index = header_index.get(column_name)
            if index is None or index >= len(row):
                return ''
            return row[index]

        inventory_number = normalize_excel_value(get_value('inventory_number'))
        asset_type_name = normalize_excel_value(get_value('asset_type'))
        status_name = normalize_excel_value(get_value('status'))

        if not inventory_number or not asset_type_name or not status_name:
            skipped_rows.append(f'Строка {row_number}: обязательные поля inventory_number, asset_type и status должны быть заполнены')
            continue

        employee_username = normalize_excel_value(get_value('employee_username'))
        try:
            purchase_date = parse_excel_date(get_value('purchase_date'))
            price = parse_excel_price(get_value('price'))
        except ValueError as exc:
            skipped_rows.append(f'Строка {row_number}: {exc}')
            continue

        employee = None
        if employee_username:
            employee = User.objects.filter(username=employee_username).first()
            if employee is None:
                skipped_rows.append(f'Строка {row_number}: пользователь {employee_username} не найден')
                continue

        asset_type, _ = AssetType.objects.get_or_create(name=asset_type_name)
        status, _ = Status.objects.get_or_create(name=status_name)

        location_name = normalize_excel_value(get_value('location'))
        location = None
        if location_name:
            location, _ = Location.objects.get_or_create(name=location_name)

        asset_defaults = {
            'asset_name': normalize_excel_value(get_value('asset_name')),
            'serial_number': normalize_excel_value(get_value('serial_number')) or None,
            'asset_type': asset_type,
            'status': status,
            'employee': employee,
            'location': location,
            'purchase_date': purchase_date,
            'price': price,
            'description': normalize_excel_value(get_value('description')),
        }

        asset, created = Asset.objects.update_or_create(
            inventory_number=inventory_number,
            defaults=asset_defaults,
        )

        if created:
            created_count += 1
        else:
            updated_count += 1

    return created_count, updated_count, skipped_rows


def require_view_access(profile, section_name, message="Недостаточно прав на просмотр раздела"):
    if not profile.has_section_access(section_name):
        return deny_access(message)
    return None


def require_full_access(profile, section_name, message="Недостаточно прав для выполнения действия"):
    if not profile.has_full_section_access(section_name):
        return deny_access(message)
    return None


def can_manage_target_employee(profile, target_user):
    if profile.can_manage_privileged_employees():
        return True

    if is_privileged_user(target_user):
        return False

    return is_employee_or_unassigned_user(target_user)


def redirect_employees_access_denied(request, message="Нет доступа к выбранному сотруднику"):
    messages.error(request, message)
    return redirect('employees')

# Для экспорта в Excel
try:
    from openpyxl import Workbook
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# Для экспорта в PDF
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os
    
    # Регистрируем шрифт с поддержкой кириллицы (Arial)
    # Сначала пробуем системные пути к шрифтам
    font_paths = [
        'C:/Windows/Fonts/arial.ttf',           # Windows
        'C:/Windows/Fonts/arialbd.ttf',          # Windows Bold
        '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',  # Linux
        '/usr/share/fonts/truetype/ubuntu/Ubuntu-R.ttf',                    # Linux Ubuntu
    ]
    
    font_registered = False
    for font_path in font_paths:
        if os.path.exists(font_path):
            pdfmetrics.registerFont(TTFont('Arial', font_path))
            font_registered = True
            break
    
    if not font_registered:
        # Если системные шрифты не найдены, используем встроенный Helvetica
        # но предупреждаем что кириллица может не отображаться
        pass
    
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')
    
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)
            # Логируем авторизацию
            log_action(user, ActionLog.ActionType.LOGIN, f'Пользователь {user.username} авторизовался')
            return redirect('home')
    else:
        form = AuthenticationForm()
    
    return render(request, 'login.html', {'form': form})


def logout_view(request):
    user = request.user
    if user.is_authenticated:
        log_action(user, ActionLog.ActionType.LOGOUT, f'Пользователь {user.username} вышел из системы')
    logout(request)
    return redirect('login')


@login_required
def home(request):
    profile = get_user_access(request.user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()
    is_employee = profile.is_employee()
    
    # Получаем последние действия для администраторов
    recent_actions = []
    if is_admin:
        recent_actions = ActionLog.objects.select_related('user').all().order_by('-created_at')[:50]
    
    return render(request, 'home.html', {
        'is_admin': is_admin, 
        'is_manager': is_manager,
        'is_employee': is_employee,
        'profile': profile,
        'recent_actions': recent_actions
    })


@login_required
def employees_list(request):
    """Список сотрудников с фильтрацией и поиском"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'employees', "Недостаточно прав для просмотра сотрудников")
    if access_error:
        return access_error
    
    users = get_visible_employees_queryset(profile)
    
    # Поиск по имени, фамилии, username, email
    query = request.GET.get('q', '').strip()
    if query:
        users = users.filter(
            models.Q(username__icontains=query) |
            models.Q(first_name__icontains=query) |
            models.Q(last_name__icontains=query) |
            models.Q(email__icontains=query)
        )
    
    # Фильтр по роли
    group_id = request.GET.get('role')
    if group_id:
        users = users.filter(groups__id=group_id)
    
    # Фильтр по статусу (активен/неактивен)
    is_active = request.GET.get('is_active')
    if is_active == '1':
        users = users.filter(is_active=True)
    elif is_active == '0':
        users = users.filter(is_active=False)
    
    # Добавляем флаг can_edit для каждого пользователя
    users_with_permissions = []
    for u in users:
        can_edit = profile.can_manage_employees() and can_manage_target_employee(profile, u)
        primary_group_name = get_primary_group_name(u)
        role_badge_class = 'none'
        if primary_group_name == 'Администратор':
            role_badge_class = 'admin'
        elif primary_group_name == 'Руководитель':
            role_badge_class = 'manager'
        elif primary_group_name == 'Сотрудник':
            role_badge_class = 'employee'
        elif primary_group_name:
            role_badge_class = 'user'

        users_with_permissions.append({
            'user': u,
            'can_edit': can_edit,
            'role_badge_class': role_badge_class,
        })

    groups = get_manageable_groups(profile)
    
    context = {
        'users_with_permissions': users_with_permissions,
        'roles': groups,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'can_manage_employees': profile.can_manage_employees(),
    }
    return render(request, 'employees.html', context)


@login_required
def employee_add(request):
    """Добавление нового сотрудника"""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'employees', "Недостаточно прав для добавления сотрудников")
    if access_error:
        return access_error
    
    allowed_groups = get_assignable_groups(profile)
    
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        email = request.POST.get('email', '')
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        group_id = request.POST.get('role')
        target_group = None
        if group_id:
            target_group = allowed_groups.filter(id=group_id).first()
            if target_group is None:
                return render(request, 'employee_add.html', {
                    'error': 'Вы не можете назначить выбранную группу',
                    'roles': allowed_groups
                })
        
        if User.objects.filter(username=username).exists():
            return render(request, 'employee_add.html', {
                'error': 'Пользователь с таким именем уже существует',
                'roles': allowed_groups
            })
        
        user = User.objects.create_user(username=username, password=password, email=email, first_name=first_name, last_name=last_name)
        if target_group:
            user.groups.set([target_group])
        else:
            user.groups.clear()
        
        # Логируем добавление сотрудника
        log_action(request.user, ActionLog.ActionType.EMPLOYEE_CREATE, f'Добавлен сотрудник: {username}')
        
        return redirect('employees')
    
    return render(request, 'employee_add.html', {'roles': allowed_groups})


@login_required
def employee_edit(request, pk):
    """Редактирование сотрудника"""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'employees', "Недостаточно прав для редактирования сотрудников")
    if access_error:
        return access_error
    
    target_user = get_object_or_404(get_visible_employees_queryset(profile), pk=pk)

    if not can_manage_target_employee(profile, target_user):
        return redirect_employees_access_denied(request, "Нет доступа к редактированию этого пользователя")
    
    allowed_groups = get_assignable_groups(profile)
    
    if request.method == 'POST':
        target_user.first_name = request.POST.get('first_name', '')
        target_user.last_name = request.POST.get('last_name', '')
        target_user.email = request.POST.get('email', '')
        password = request.POST.get('password')
        group_id = request.POST.get('role')
        target_group = None
        if group_id:
            target_group = allowed_groups.filter(id=group_id).first()
            if target_group is None:
                return render(request, 'employee_form.html', {
                    'error': 'Вы не можете назначить выбранную группу',
                    'employee_user': target_user,
                    'roles': allowed_groups,
                    'selected_group_id': group_id,
                })
        
        if password:
            target_user.set_password(password)
        
        if target_group:
            target_user.groups.set([target_group])
        else:
            target_user.groups.clear()
        target_user.save()
        
        # Логируем редактирование сотрудника
        log_action(request.user, ActionLog.ActionType.EMPLOYEE_EDIT, f'Отредактирован сотрудник: {target_user.username}')
        
        return redirect('employees')
    
    return render(request, 'employee_form.html', {
        'employee_user': target_user,
        'roles': allowed_groups,
        'selected_group_id': get_primary_group(target_user).id if get_primary_group(target_user) else '',
    })


@login_required
def employee_delete(request, pk):
    """Удаление сотрудника"""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'employees', "Недостаточно прав для удаления сотрудников")
    if access_error:
        return access_error
    
    if request.method == 'POST':
        target_user = get_object_or_404(get_visible_employees_queryset(profile), pk=pk)

        if not can_manage_target_employee(profile, target_user):
            return redirect_employees_access_denied(request, "Нет доступа к удалению этого пользователя")
        
        if target_user != user:
            target_user.delete()
            
            # Логируем удаление сотрудника
            log_action(request.user, ActionLog.ActionType.EMPLOYEE_DELETE, f'Удалён сотрудник: {target_user.username}')
        
        return redirect('employees')
    
    return redirect('employees')


# ====== Управление активами ======

@login_required
def assets_list(request):
    """Список всех активов с фильтрацией и поиском"""
    profile = get_user_access(request.user)
    access_error = require_view_access(profile, 'assets', "Недостаточно прав для просмотра активов")
    if access_error:
        return access_error

    if request.method == 'POST':
        manage_error = require_full_access(profile, 'assets', "Недостаточно прав для согласования личных устройств")
        if manage_error:
            return manage_error

        byod_action = request.POST.get('byod_action')
        if byod_action in {'approve', 'revoke'}:
            resource = get_object_or_404(BYODResource.objects.select_related('employee'), pk=request.POST.get('resource_id'))
            resource.is_company_approved = byod_action == 'approve'
            resource.save(update_fields=['is_company_approved', 'updated_at'])

            action_label = 'согласован' if resource.is_company_approved else 'снят с согласования'
            log_action(
                request.user,
                ActionLog.ActionType.ASSET_EDIT,
                f'BYOD-ресурс {resource.device_name} ({resource.employee.username}) {action_label}',
                resource.pk,
            )
            messages.success(
                request,
                f'Статус личного устройства «{resource.device_name}» пользователя {resource.employee.get_full_name() or resource.employee.username} обновлён.'
            )
        else:
            messages.warning(request, 'Неизвестное действие для личного устройства.')

        return redirect('assets')

    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).all()
    
    # Поиск по названию, инвентарному или серийному номеру
    query = request.GET.get('q', '').strip()
    if query:
        assets = assets.filter(
            models.Q(asset_name__icontains=query) |
            models.Q(inventory_number__icontains=query) |
            models.Q(serial_number__icontains=query) |
            models.Q(description__icontains=query)
        )
    
    # Фильтр по типу актива
    asset_type_id = request.GET.get('type')
    if asset_type_id:
        assets = assets.filter(asset_type_id=asset_type_id)
    
    # Фильтр по статусу
    status_id = request.GET.get('status')
    if status_id:
        assets = assets.filter(status_id=status_id)
    
    # Фильтр по сотруднику
    employee_id = request.GET.get('employee')
    if employee_id:
        assets = assets.filter(employee_id=employee_id)
    
    # Формируем списки для выбора в фильтрах
    type_choices = [(at.id, at.name) for at in AssetType.objects.all()]
    status_choices = [(s.id, s.name) for s in Status.objects.all()]
    
    byod_resources = BYODResource.objects.select_related('employee').order_by(
        'is_company_approved', 'employee__first_name', 'employee__username', '-updated_at'
    )

    context = {
        'assets': assets,
        'type_choices': type_choices,
        'status_choices': status_choices,
        'employees': User.objects.all(),
        'byod_resources': byod_resources,
        'byod_total_count': byod_resources.count(),
        'byod_approved_count': byod_resources.filter(is_company_approved=True).count(),
        'byod_pending_count': byod_resources.filter(is_company_approved=False).count(),
        'can_manage_assets': profile.can_manage_assets(),
        'can_view_asset_card': profile.can_view_asset_card(),
        'excel_import_available': OPENPYXL_AVAILABLE,
    }
    return render(request, 'assets.html', context)


@login_required
def inventory_list(request):
    """Минимальный раздел инвентаризации корпоративных активов."""
    profile = get_user_access(request.user)
    access_error = require_view_access(profile, 'assets', "Недостаточно прав для просмотра инвентаризации")
    if access_error:
        return access_error

    if request.method == 'POST':
        manage_error = require_full_access(profile, 'assets', "Недостаточно прав для создания инвентаризации")
        if manage_error:
            return manage_error

        title = request.POST.get('title', '').strip()
        location_id = request.POST.get('location') or None
        notes = request.POST.get('notes', '').strip()

        if not title:
            messages.error(request, 'Укажите название инвентаризации.')
            return redirect('inventory_list')

        location = None
        if location_id:
            location = get_object_or_404(Location, pk=location_id)

        session = InventorySession.objects.create(
            title=title,
            location=location,
            created_by=request.user,
            started_at=timezone.now(),
            status=InventorySession.StatusChoices.IN_PROGRESS,
            notes=notes,
        )
        build_inventory_items_for_session(session)
        log_action(
            request.user,
            ActionLog.ActionType.INVENTORY_CREATE,
            f'Создана инвентаризация: {session.title}',
            session.pk,
        )
        messages.success(request, f'Инвентаризация «{session.title}» создана.')
        return redirect('inventory_detail', pk=session.pk)

    sessions = InventorySession.objects.select_related('location', 'created_by').prefetch_related('items').all()

    session_query = request.GET.get('q', '').strip()
    selected_status = request.GET.get('status', '').strip()
    selected_location = request.GET.get('location', '').strip()

    if session_query:
        sessions = sessions.filter(
            Q(title__icontains=session_query)
            | Q(notes__icontains=session_query)
            | Q(created_by__username__icontains=session_query)
            | Q(created_by__first_name__icontains=session_query)
            | Q(created_by__last_name__icontains=session_query)
            | Q(location__name__icontains=session_query)
        )

    if selected_status:
        sessions = sessions.filter(status=selected_status)

    if selected_location:
        sessions = sessions.filter(location_id=selected_location)

    status_choices = InventorySession.StatusChoices.choices
    session_stats = []
    for session in sessions:
        total = session.items.count()
        checked = session.items.exclude(result=InventoryItem.ResultChoices.PENDING).count()
        mismatches = session.items.filter(result=InventoryItem.ResultChoices.MISMATCH).count()
        missing = session.items.filter(result=InventoryItem.ResultChoices.MISSING).count()
        session_stats.append({
            'session': session,
            'total': total,
            'checked': checked,
            'pending': max(total - checked, 0),
            'mismatches': mismatches,
            'missing': missing,
        })

    context = {
        'session_stats': session_stats,
        'locations': Location.objects.order_by('name'),
        'status_choices': status_choices,
        'selected_query': session_query,
        'selected_status': selected_status,
        'selected_location': selected_location,
        'can_manage_assets': profile.can_manage_assets(),
    }
    return render(request, 'inventory_list.html', context)


@login_required
def inventory_detail(request, pk):
    """Карточка инвентаризации и отметка результатов по активам."""
    profile = get_user_access(request.user)
    access_error = require_view_access(profile, 'assets', "Недостаточно прав для просмотра инвентаризации")
    if access_error:
        return access_error

    session = get_object_or_404(
        InventorySession.objects.select_related('location', 'created_by'),
        pk=pk,
    )

    if request.method == 'POST':
        manage_error = require_full_access(profile, 'assets', "Недостаточно прав для проведения инвентаризации")
        if manage_error:
            return manage_error

        item = get_object_or_404(
            InventoryItem.objects.select_related('asset', 'expected_employee', 'expected_location', 'expected_status'),
            pk=request.POST.get('item_id'),
            session=session,
        )
        result = request.POST.get('result', InventoryItem.ResultChoices.PENDING)
        comment = request.POST.get('comment', '').strip()
        actual_employee_id = request.POST.get('actual_employee') or None
        actual_location_id = request.POST.get('actual_location') or None
        actual_status_id = request.POST.get('actual_status') or None

        valid_results = {choice[0] for choice in InventoryItem.ResultChoices.choices}
        if result not in valid_results or result == InventoryItem.ResultChoices.PENDING:
            messages.error(request, 'Выберите корректный результат проверки.')
            return redirect('inventory_detail', pk=session.pk)

        item.result = result
        item.comment = comment
        item.checked_by = request.user
        item.checked_at = timezone.now()
        item.actual_employee_id = actual_employee_id
        item.actual_location_id = actual_location_id
        item.actual_status_id = actual_status_id

        if result == InventoryItem.ResultChoices.MATCHED:
            item.actual_employee = item.expected_employee
            item.actual_location = item.expected_location
            item.actual_status = item.expected_status

        item.save()
        log_action(
            request.user,
            ActionLog.ActionType.INVENTORY_CHECK,
            f'Отмечен актив {item.asset.inventory_number} в инвентаризации {session.title}: {item.get_result_display()}',
            item.asset_id,
        )
        messages.success(request, f'Результат по активу {item.asset.inventory_number} сохранён.')
        return redirect('inventory_detail', pk=session.pk)

    items = session.items.select_related(
        'asset', 'expected_employee', 'expected_location', 'expected_status',
        'actual_employee', 'actual_location', 'actual_status', 'checked_by'
    ).all()
    result_filter = request.GET.get('result', '').strip()
    if result_filter:
        items = items.filter(result=result_filter)

    total = session.items.count()
    checked = session.items.exclude(result=InventoryItem.ResultChoices.PENDING).count()
    context = {
        'session': session,
        'items': items,
        'result_choices': InventoryItem.ResultChoices.choices,
        'employees': User.objects.order_by('first_name', 'username'),
        'locations': Location.objects.order_by('name'),
        'statuses': Status.objects.order_by('name'),
        'can_manage_assets': profile.can_manage_assets(),
        'total_count': total,
        'checked_count': checked,
        'pending_count': max(total - checked, 0),
        'matched_count': session.items.filter(result=InventoryItem.ResultChoices.MATCHED).count(),
        'mismatch_count': session.items.filter(result=InventoryItem.ResultChoices.MISMATCH).count(),
        'missing_count': session.items.filter(result=InventoryItem.ResultChoices.MISSING).count(),
        'selected_result': result_filter,
        'pending_value': InventoryItem.ResultChoices.PENDING,
        'matched_value': InventoryItem.ResultChoices.MATCHED,
        'mismatch_value': InventoryItem.ResultChoices.MISMATCH,
        'missing_value': InventoryItem.ResultChoices.MISSING,
    }
    return render(request, 'inventory_detail.html', context)


@login_required
def inventory_complete(request, pk):
    """Завершение инвентаризации."""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для завершения инвентаризации")
    if access_error:
        return access_error

    session = get_object_or_404(InventorySession, pk=pk)
    if request.method == 'POST':
        session.status = InventorySession.StatusChoices.COMPLETED
        session.completed_at = timezone.now()
        session.save(update_fields=['status', 'completed_at', 'updated_at'])
        log_action(
            request.user,
            ActionLog.ActionType.INVENTORY_COMPLETE,
            f'Завершена инвентаризация: {session.title}',
            session.pk,
        )
        messages.success(request, f'Инвентаризация «{session.title}» завершена.')
    return redirect('inventory_detail', pk=session.pk)


@login_required
def asset_import_template(request):
    """Скачать шаблон Excel для импорта активов"""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для скачивания шаблона импорта")
    if access_error:
        return access_error

    if not OPENPYXL_AVAILABLE:
        return HttpResponse('OpenPyXL не установлен', status=500)

    workbook = build_asset_import_workbook()
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="asset_import_template.xlsx"'
    workbook.save(response)
    return response


@login_required
def asset_import_excel(request):
    """Импорт активов из Excel файла"""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для импорта активов")
    if access_error:
        return access_error

    if request.method != 'POST':
        return redirect('assets')

    if not OPENPYXL_AVAILABLE:
        messages.error(request, 'Импорт невозможен: библиотека OpenPyXL не установлена')
        return redirect('assets')

    excel_file = request.FILES.get('excel_file')
    if not excel_file:
        messages.error(request, 'Выберите Excel файл для импорта')
        return redirect('assets')

    if not excel_file.name.lower().endswith('.xlsx'):
        messages.error(request, 'Поддерживаются только файлы формата .xlsx')
        return redirect('assets')

    try:
        workbook = load_workbook(excel_file)
        created_count, updated_count, skipped_rows = import_assets_from_workbook(workbook)
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect('assets')
    except Exception as exc:
        messages.error(request, f'Ошибка при чтении Excel файла: {exc}')
        return redirect('assets')

    if created_count or updated_count:
        messages.success(
            request,
            f'Импорт завершён: создано {created_count}, обновлено {updated_count}'
        )
        log_action(
            request.user,
            ActionLog.ActionType.ASSET_CREATE,
            f'Импорт активов из Excel: создано {created_count}, обновлено {updated_count}'
        )
    else:
        messages.warning(request, 'Импорт завершён без изменений')

    for skipped_row in skipped_rows[:10]:
        messages.warning(request, skipped_row)
    if len(skipped_rows) > 10:
        messages.warning(request, f'И ещё {len(skipped_rows) - 10} строк пропущено')

    return redirect('assets')


@login_required
def asset_add(request):
    """Добавление нового актива"""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для добавления активов")
    if access_error:
        return access_error
    
    if request.method == 'POST':
        inventory_number = request.POST.get('inventory_number')
        
        if Asset.objects.filter(inventory_number=inventory_number).exists():
            return render(request, 'asset_form.html', {
                'error': 'Актив с таким инвентарным номером уже существует',
                'asset_types': AssetType.objects.all(),
                'statuses': Status.objects.all(),
                'employees': User.objects.all(),
                'locations': Location.objects.all(),
            })
        
        asset = Asset.objects.create(
            asset_name=request.POST.get('asset_name') or '',
            inventory_number=inventory_number,
            serial_number=request.POST.get('serial_number') or None,
            asset_type_id=request.POST.get('asset_type') or None,
            status_id=request.POST.get('status') or None,
            employee_id=request.POST.get('employee') or None,
            location_id=request.POST.get('location') or None,
            purchase_date=request.POST.get('purchase_date') or None,
            price=request.POST.get('price') or None,
            description=request.POST.get('description', ''),
        )
        
        # Логируем создание актива
        log_action(request.user, ActionLog.ActionType.ASSET_CREATE, f'Добавлен актив: {inventory_number}', asset.pk)
        
        return redirect('assets')
    
    return render(request, 'asset_form.html', {
        'asset_types': AssetType.objects.all(),
        'statuses': Status.objects.all(),
        'employees': User.objects.all(),
        'locations': Location.objects.all(),
    })


@login_required
def asset_edit(request, pk):
    """Редактирование актива"""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для редактирования активов")
    if access_error:
        return access_error
    
    asset = get_object_or_404(Asset, pk=pk)
    
    if request.method == 'POST':
        inventory_number = request.POST.get('inventory_number')
        
        if Asset.objects.filter(inventory_number=inventory_number).exclude(pk=pk).exists():
            return render(request, 'asset_form.html', {
                'error': 'Актив с таким инвентарным номером уже существует',
                'asset': asset,
                'asset_types': AssetType.objects.all(),
                'statuses': Status.objects.all(),
                'employees': User.objects.all(),
                'locations': Location.objects.all(),
            })
        
        asset.asset_name = request.POST.get('asset_name') or ''
        asset.inventory_number = inventory_number
        asset.serial_number = request.POST.get('serial_number') or None
        asset.asset_type_id = request.POST.get('asset_type') or None
        asset.status_id = request.POST.get('status') or None
        asset.employee_id = request.POST.get('employee') or None
        asset.location_id = request.POST.get('location') or None
        asset.purchase_date = request.POST.get('purchase_date') or None
        asset.price = request.POST.get('price') or None
        asset.description = request.POST.get('description', '')
        asset.save()
        
        # Логируем редактирование актива
        log_action(request.user, ActionLog.ActionType.ASSET_EDIT, f'Отредактирован актив: {asset.inventory_number}', asset.pk)
        
        return redirect('assets')
    
    return render(request, 'asset_form.html', {
        'asset': asset,
        'asset_types': AssetType.objects.all(),
        'statuses': Status.objects.all(),
        'employees': User.objects.all(),
        'locations': Location.objects.all(),
    })


@login_required
def asset_delete(request, pk):
    """Удаление актива"""
    profile = get_user_access(request.user)
    access_error = require_full_access(profile, 'assets', "Недостаточно прав для удаления активов")
    if access_error:
        return access_error
    
    if request.method == 'POST':
        asset = get_object_or_404(Asset, pk=pk)
        inventory_number = asset.inventory_number
        asset.delete()
        
        # Логируем удаление актива
        log_action(request.user, ActionLog.ActionType.ASSET_DELETE, f'Удалён актив: {inventory_number}')
        
        return redirect('assets')
    
    return redirect('assets')


# ====== Мои активы ======

@login_required
def my_assets(request):
    """Список активов текущего пользователя со статусом 'В эксплуатации'"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'my_assets', "Недостаточно прав для просмотра раздела 'Мои активы'")
    if access_error:
        return access_error

    if request.method == 'POST':
        action = request.POST.get('byod_action')

        if action == 'confirm':
            resource = get_object_or_404(BYODResource, pk=request.POST.get('resource_id'), employee=user)
            resource.employee_consent = True
            resource.last_confirmed_at = timezone.now()
            resource.save(update_fields=['employee_consent', 'last_confirmed_at', 'updated_at'])
            log_action(user, ActionLog.ActionType.ASSET_EDIT, f'Подтверждён BYOD-ресурс {resource.device_name}', resource.pk)
            messages.success(request, 'Согласие по личному устройству подтверждено.')
            return redirect('my_assets')

        if action == 'delete':
            resource = get_object_or_404(BYODResource, pk=request.POST.get('resource_id'), employee=user)
            resource_name = resource.device_name
            resource.delete()
            log_action(user, ActionLog.ActionType.ASSET_DELETE, f'Удалён BYOD-ресурс {resource_name}')
            messages.success(request, f'Личное устройство «{resource_name}» удалено.')
            return redirect('my_assets')

        if action in {'add', 'edit'}:
            device_name = request.POST.get('device_name', '').strip()
            resource_type = request.POST.get('resource_type', BYODResource.ResourceType.OTHER)
            brand_model = request.POST.get('brand_model', '').strip()
            serial_number = request.POST.get('serial_number', '').strip()
            usage_status = request.POST.get('usage_status', BYODResource.UsageStatus.ACTIVE)
            work_purpose = request.POST.get('work_purpose', '').strip()
            support_notes = request.POST.get('support_notes', '').strip()
            compensation_notes = request.POST.get('compensation_notes', '').strip()
            employee_consent = request.POST.get('employee_consent') == 'on'
            requested_company_approval = request.POST.get('is_company_approved') == 'on'

            if not device_name or not work_purpose:
                messages.error(request, 'Заполните обязательные поля: название устройства и рабочее назначение.')
                return redirect('my_assets')

            if not employee_consent:
                messages.error(request, 'Нужно подтвердить согласие на мягкий учёт личного устройства.')
                return redirect('my_assets')

            if action == 'add':
                byod_resource = BYODResource.objects.create(
                    employee=user,
                    device_name=device_name,
                    resource_type=resource_type,
                    brand_model=brand_model,
                    serial_number=serial_number,
                    usage_status=usage_status,
                    work_purpose=work_purpose,
                    support_notes=support_notes,
                    compensation_notes=compensation_notes,
                    employee_consent=employee_consent,
                    is_company_approved=requested_company_approval if (is_admin or is_manager) else False,
                    last_confirmed_at=timezone.now(),
                )
                log_action(user, ActionLog.ActionType.ASSET_CREATE, f'Добавлен BYOD-ресурс {byod_resource.device_name}', byod_resource.pk)
                messages.success(request, 'Личное устройство добавлено в список ваших активов.')
                return redirect('my_assets')

            resource = get_object_or_404(BYODResource, pk=request.POST.get('resource_id'), employee=user)
            resource.device_name = device_name
            resource.resource_type = resource_type
            resource.brand_model = brand_model
            resource.serial_number = serial_number
            resource.usage_status = usage_status
            resource.work_purpose = work_purpose
            resource.support_notes = support_notes
            resource.compensation_notes = compensation_notes
            resource.employee_consent = employee_consent
            if is_admin or is_manager:
                resource.is_company_approved = requested_company_approval
            resource.last_confirmed_at = timezone.now()
            resource.save()
            log_action(user, ActionLog.ActionType.ASSET_EDIT, f'Обновлён BYOD-ресурс {resource.device_name}', resource.pk)
            messages.success(request, f'Личное устройство «{resource.device_name}» обновлено.')
            return redirect('my_assets')
    
    # Получаем статус "В эксплуатации"
    try:
        in_operation_status = Status.objects.get(name='В эксплуатации')
        assets = Asset.objects.select_related(
            'asset_type', 'status', 'employee', 'location'
        ).filter(
            employee=user,
            status=in_operation_status
        )
    except Status.DoesNotExist:
        assets = Asset.objects.none()

    byod_resources = BYODResource.objects.filter(employee=user).order_by('-updated_at', '-created_at')
    
    context = {
        'assets': assets,
        'byod_resources': byod_resources,
        'byod_resources_count': byod_resources.count(),
        'resource_types': BYODResource.ResourceType.choices,
        'usage_statuses': BYODResource.UsageStatus.choices,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'can_view_asset_card': profile.can_view_asset_card(),
    }
    return render(request, 'my_assets.html', context)


@login_required
def my_byod_resources(request):
    """Учёт личной техники сотрудника, используемой в работе."""
    return redirect('my_assets')


# ====== Мои запросы ======

@login_required
def my_requests(request):
    """Список запросов на ремонт текущего пользователя"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'my_requests', "Недостаточно прав для просмотра раздела 'Мои запросы'")
    if access_error:
        return access_error
    
    # Получаем все запросы текущего пользователя
    repair_requests = RepairRequest.objects.select_related(
        'asset', 'asset__asset_type', 'asset__status', 'asset__location',
        'processed_by'
    ).filter(requester=user).order_by('-created_at')
    
    # Фильтр по типу актива
    asset_type_id = request.GET.get('asset_type')
    selected_asset_type = None
    if asset_type_id:
        repair_requests = repair_requests.filter(asset__asset_type_id=asset_type_id)
        selected_asset_type = AssetType.objects.filter(id=asset_type_id).first()
    
    # Фильтр по статусу запроса
    selected_status = request.GET.get('status')
    if selected_status:
        repair_requests = repair_requests.filter(status=selected_status)
    
    # Получаем все типы активов для фильтра
    asset_types = AssetType.objects.all()
    
    context = {
        'repair_requests': repair_requests,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'asset_types': asset_types,
        'selected_asset_type': selected_asset_type,
        'selected_status': selected_status,
    }
    return render(request, 'my_requests.html', context)


# ====== Запросы на ремонт ======

@login_required
def create_repair_request(request, asset_id):
    """Создание запроса на ремонт актива"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_full_access(profile, 'my_assets', "Недостаточно прав для создания запроса на ремонт")
    if access_error:
        return access_error
    
    asset = get_object_or_404(Asset, pk=asset_id)
    
    # Проверяем, что актив принадлежит пользователю и имеет статус "В эксплуатации"
    if asset.employee != user and not is_admin and not is_manager:
        return HttpResponseForbidden("Вы не можете запросить ремонт этого актива")
    
    try:
        in_operation_status = Status.objects.get(name='В эксплуатации')
        if asset.status != in_operation_status:
            return HttpResponseForbidden("Актив должен иметь статус 'В эксплуатации'")
    except Status.DoesNotExist:
        return HttpResponseForbidden("Статус 'В эксплуатации' не найден")
    
    # Проверяем, нет ли уже активного запроса
    existing_request = RepairRequest.objects.filter(
        asset=asset,
        status=RepairRequest.StatusChoices.PENDING
    ).exists()
    
    if existing_request:
        return HttpResponseForbidden("Запрос на ремонт этого актива уже существует")
    
    reason = request.POST.get('reason', '') if request.method == 'POST' else ''
    
    # Создаём запрос на ремонт
    repair_req = RepairRequest.objects.create(
        asset=asset,
        requester=user,
        status=RepairRequest.StatusChoices.PENDING,
        reason=reason
    )
    
    # Логируем создание запроса на ремонт
    log_action(user, ActionLog.ActionType.REQUEST_CREATE, f'Создан запрос на ремонт актива {asset.inventory_number}', repair_req.pk)
    
    return redirect('my_assets')


@login_required
def repair_requests_list(request):
    """Список запросов на ремонт для администраторов и руководителей"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'repair_requests', "Недостаточно прав для просмотра запросов на ремонт")
    if access_error:
        return access_error
    
    # Получаем только ожидающие запросы
    repair_requests = RepairRequest.objects.select_related(
        'asset', 'asset__asset_type', 'asset__status', 'asset__employee', 
        'asset__location', 'requester'
    ).filter(status=RepairRequest.StatusChoices.PENDING)
    
    # Фильтр по типу актива
    asset_type_id = request.GET.get('asset_type')
    if asset_type_id:
        repair_requests = repair_requests.filter(asset__asset_type_id=asset_type_id)
    
    # Фильтр по сотруднику
    employee_id = request.GET.get('employee')
    if employee_id:
        repair_requests = repair_requests.filter(asset__employee_id=employee_id)
    
    # Фильтр по локации
    location_id = request.GET.get('location')
    if location_id:
        repair_requests = repair_requests.filter(asset__location_id=location_id)
    
    # Получаем данные для фильтров
    asset_types = AssetType.objects.all()
    employees = User.objects.exclude(groups__permissions__codename='manage_privileged_employees').distinct()
    locations = Location.objects.all()
    
    # Получаем выбранные значения фильтров
    selected_asset_type = AssetType.objects.filter(id=asset_type_id).first() if asset_type_id else None
    selected_employee = User.objects.filter(id=employee_id).first() if employee_id else None
    selected_location = Location.objects.filter(id=location_id).first() if location_id else None
    
    context = {
        'repair_requests': repair_requests,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'asset_types': asset_types,
        'employees': employees,
        'locations': locations,
        'selected_asset_type': selected_asset_type,
        'selected_employee': selected_employee,
        'selected_location': selected_location,
        'can_manage_repair_requests': profile.can_manage_repair_requests(),
    }
    return render(request, 'repair_requests.html', context)


@login_required
def approve_repair_request(request, pk):
    """Принятие запроса на ремонт"""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'repair_requests', "Недостаточно прав для обработки запросов на ремонт")
    if access_error:
        return access_error
    
    repair_request = get_object_or_404(RepairRequest, pk=pk)
    
    if repair_request.status != RepairRequest.StatusChoices.PENDING:
        return HttpResponseForbidden("Запрос уже обработан")
    
    if request.method == 'POST':
        # Получаем статус "На ремонте"
        try:
            repair_status = Status.objects.get(name='На ремонте')
            repair_request.asset.status = repair_status
            repair_request.asset.save()
        except Status.DoesNotExist:
            pass  # Если статус не найден, просто обновляем запрос
        
        # Обновляем статус запроса
        repair_request.status = RepairRequest.StatusChoices.APPROVED
        repair_request.processed_by = user
        repair_request.processed_at = timezone.now()
        repair_request.save()
        
        # Логируем принятие запроса на ремонт
        log_action(user, ActionLog.ActionType.REQUEST_APPROVE, f'Принят запрос на ремонт актива {repair_request.asset.inventory_number}', repair_request.pk)
        
        return redirect('repair_requests')
    
    return redirect('repair_requests')


@login_required
def reject_repair_request(request, pk):
    """Отклонение запроса на ремонт"""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'repair_requests', "Недостаточно прав для обработки запросов на ремонт")
    if access_error:
        return access_error
    
    repair_request = get_object_or_404(RepairRequest, pk=pk)
    
    if repair_request.status != RepairRequest.StatusChoices.PENDING:
        return HttpResponseForbidden("Запрос уже обработан")
    
    if request.method == 'POST':
        repair_request.status = RepairRequest.StatusChoices.REJECTED
        repair_request.processed_by = user
        repair_request.processed_at = timezone.now()
        repair_request.save()
        
        # Логируем отклонение запроса на ремонт
        log_action(user, ActionLog.ActionType.REQUEST_REJECT, f'Отклонён запрос на ремонт актива {repair_request.asset.inventory_number}', repair_request.pk)
        
        return redirect('repair_requests')
    
    return redirect('repair_requests')


# ====== QR-код актива ======

@login_required
def asset_qr(request, pk):
    """Генерация QR-кода для актива"""
    profile = get_user_access(request.user)
    access_error = require_view_access(profile, 'assets', "Недостаточно прав для просмотра QR-кода актива")
    if access_error:
        return access_error

    asset = get_object_or_404(Asset, pk=pk)
    
    # Формируем ссылку для QR-кода
    base_url = 'http://127.0.0.1:8000'
    qr_url = f"{base_url}/assets/{asset.pk}"
    
    # Генерируем QR-код
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=4,
    )
    qr.add_data(qr_url)
    qr.make(fit=True)
    
    img = qr.make_image(fill_color="black", back_color="white")
    
    # Конвертируем в base64 для отображения
    buffer = BytesIO()
    img.save(buffer, format='PNG')
    image_base64 = base64.b64encode(buffer.getvalue()).decode()
    
    context = {
        'asset': asset,
        'qr_url': qr_url,
        'qr_image': image_base64,
    }
    return render(request, 'asset_qr.html', context)


@login_required
def asset_detail(request, pk):
    """Просмотр информации об активе"""
    profile = get_user_access(request.user)
    access_error = require_permission(profile, 'can_view_asset_card', "Недостаточно прав для просмотра карточки актива")
    if access_error:
        return access_error

    asset = get_object_or_404(Asset, pk=pk)

    is_admin = profile.is_administrator()
    can_transfer = can_transfer_asset(profile, asset)

    if request.method == 'POST':
        if not can_transfer:
            return HttpResponseForbidden("Недостаточно прав для передачи актива")

        transfer_to_user_id = request.POST.get('transfer_to_user')
        if not transfer_to_user_id:
            messages.error(request, 'Выберите пользователя, которому нужно передать актив')
            return redirect('asset_detail', pk=asset.pk)

        new_owner = User.objects.filter(pk=transfer_to_user_id).first()
        if new_owner is None:
            messages.error(request, 'Выбранный пользователь не найден')
            return redirect('asset_detail', pk=asset.pk)

        if asset.employee_id == new_owner.id:
            messages.error(request, 'Нельзя передать актив текущему владельцу')
            return redirect('asset_detail', pk=asset.pk)

        previous_owner = asset.employee
        asset.employee = new_owner
        asset.save(update_fields=['employee', 'updated_at'])

        previous_owner_name = previous_owner.get_full_name().strip() if previous_owner else ''
        new_owner_name = new_owner.get_full_name().strip()
        previous_owner_label = previous_owner_name or (previous_owner.username if previous_owner else 'Не назначен')
        new_owner_label = new_owner_name or new_owner.username

        log_action(
            request.user,
            ActionLog.ActionType.ASSET_TRANSFER,
            f'Актив {asset.inventory_number} передан от {previous_owner_label} пользователю {new_owner_label}',
            asset.pk,
        )
        messages.success(request, f'Актив {asset.inventory_number} успешно передан пользователю {new_owner_label}')
        return redirect('asset_detail', pk=asset.pk)

    transferable_users = User.objects.filter(is_active=True).order_by('last_name', 'first_name', 'username') if can_transfer else User.objects.none()
    
    context = {
        'asset': asset,
        'is_admin': is_admin,
        'can_manage_assets': profile.can_manage_assets(),
        'can_transfer_asset': can_transfer,
        'can_view_asset_history': profile.can_view_asset_history(),
        'transferable_users': transferable_users,
    }
    return render(request, 'asset_detail.html', context)


@login_required
def my_byod_resource_detail(request, pk):
    """Просмотр информации о личном активе сотрудника."""
    profile = get_user_access(request.user)
    resource = get_object_or_404(
        BYODResource.objects.select_related('employee'),
        pk=pk,
    )

    can_manage_assets = profile.can_manage_assets()
    is_owner = resource.employee_id == request.user.id

    if not (is_owner or can_manage_assets):
        return deny_access("Недостаточно прав для просмотра карточки личного актива")

    from_page = request.GET.get('from')
    back_url = reverse('my_assets')
    if from_page == 'assets' and can_manage_assets:
        back_url = reverse('assets')

    context = {
        'resource': resource,
        'can_manage_assets': can_manage_assets,
        'is_owner': is_owner,
        'back_url': back_url,
        'back_to_assets': from_page == 'assets' and can_manage_assets,
    }
    return render(request, 'my_byod_resource_detail.html', context)


@login_required
def asset_history(request, pk):
    """История действий, связанных с конкретным активом"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_permission(profile, 'can_view_asset_history', "Недостаточно прав для просмотра истории актива")
    if access_error:
        return access_error

    asset = get_object_or_404(Asset, pk=pk)

    asset_action_types = [
        ActionLog.ActionType.ASSET_CREATE,
        ActionLog.ActionType.ASSET_EDIT,
        ActionLog.ActionType.ASSET_DELETE,
        ActionLog.ActionType.ASSET_TRANSFER,
    ]
    repair_action_types = [
        ActionLog.ActionType.REQUEST_CREATE,
        ActionLog.ActionType.REQUEST_APPROVE,
        ActionLog.ActionType.REQUEST_REJECT,
    ]

    actions = ActionLog.objects.select_related('user').filter(
        Q(action_type__in=asset_action_types, object_id=asset.pk)
        | Q(action_type__in=repair_action_types, description__icontains=asset.inventory_number)
    ).order_by('-created_at')

    context = {
        'asset': asset,
        'actions': actions,
        'is_admin': is_admin,
        'is_manager': is_manager,
    }
    return render(request, 'asset_history.html', context)


# ====== Отчёты ======

@login_required
def reports_list(request):
    """Список доступных отчётов"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'reports', "Недостаточно прав для просмотра отчётов")
    if access_error:
        return access_error
    
    context = {
        'is_admin': is_admin,
        'is_manager': is_manager,
    }
    return render(request, 'reports.html', context)


@login_required
def action_log(request):
    """Журнал действий пользователей в системе"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'action_log', "Недостаточно прав для просмотра журнала действий")
    if access_error:
        return access_error
    
    # Получаем все действия с пагинацией
    from django.core.paginator import Paginator
    
    actions = ActionLog.objects.select_related('user').all().order_by('-created_at')
    
    # Фильтр по типу действия
    action_type = request.GET.get('action_type')
    if action_type:
        actions = actions.filter(action_type=action_type)
    
    # Фильтр по пользователю
    user_id = request.GET.get('user')
    if user_id:
        actions = actions.filter(user_id=user_id)
    
    # Фильтр по дате
    date_from = request.GET.get('date_from')
    if date_from:
        actions = actions.filter(created_at__date__gte=date_from)
    
    date_to = request.GET.get('date_to')
    if date_to:
        actions = actions.filter(created_at__date__lte=date_to)
    
    # Пагинация
    paginator = Paginator(actions, 50)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    context = {
        'page_obj': page_obj,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'action_types': ActionLog.ActionType.choices,
        'users': User.objects.all().order_by('username'),
        'selected_action_type': action_type,
        'selected_user': user_id,
        'date_from': date_from,
        'date_to': date_to,
    }
    return render(request, 'action_log.html', context)


@login_required
def report_assets_by_employee(request):
    """Отчёт 'Какие активы у сотрудников'"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'reports', "Недостаточно прав для просмотра отчётов")
    if access_error:
        return access_error
    
    # Получаем все активы с информацией о сотрудниках
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(
        employee__isnull=False
    ).order_by('employee__last_name', 'employee__first_name', 'asset_name')
    
    context = {
        'report_title': 'Какие активы у сотрудников',
        'assets': assets,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'can_manage_reports': profile.can_manage_reports(),
    }
    
    # Логируем создание отчёта
    log_action(user, ActionLog.ActionType.REPORT_CREATE, 'Создан отчёт: Какие активы у сотрудников')
    
    return render(request, 'report_assets_by_employee.html', context)


@login_required
def report_assets_by_department(request):
    """Отчёт 'Количество активов по отделам'"""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'reports', "Недостаточно прав для просмотра отчётов")
    if access_error:
        return access_error
    
    # Получаем все активы с информацией о сотрудниках и их местоположении (отделе)
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(
        employee__isnull=False,
        location__isnull=False
    ).order_by('location__name', 'asset_type__name', 'status__name')
    
    # Формируем данные для отчёта: группируем по отделу, типу и статусу
    report_data = {}
    for asset in assets:
        location_name = asset.location.name if asset.location else 'Без локации'
        type_name = asset.asset_type.name if asset.asset_type else 'Без типа'
        status_name = asset.status.name if asset.status else 'Без статуса'
        
        if location_name not in report_data:
            report_data[location_name] = {}
        if type_name not in report_data[location_name]:
            report_data[location_name][type_name] = {}
        if status_name not in report_data[location_name][type_name]:
            report_data[location_name][type_name][status_name] = 0
        report_data[location_name][type_name][status_name] += 1
    
    # Получаем все статусы для отображения в таблице
    all_statuses = list(Status.objects.values_list('name', flat=True).order_by('name'))
    
    context = {
        'report_title': 'Количество активов по отделам',
        'report_data': report_data,
        'all_statuses': all_statuses,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'can_manage_reports': profile.can_manage_reports(),
    }
    
    # Логируем создание отчёта
    log_action(user, ActionLog.ActionType.REPORT_CREATE, 'Создан отчёт: Количество активов по отделам')
    
    return render(request, 'report_assets_by_department.html', context)


@login_required
def report_inventory_session(request):
    """Отчёт по выбранной инвентаризации с фильтрацией."""
    user = request.user
    profile = get_user_access(user)
    is_admin = profile.is_administrator()
    is_manager = profile.is_manager()

    access_error = require_view_access(profile, 'reports', "Недостаточно прав для просмотра отчётов")
    if access_error:
        return access_error

    session_search = request.GET.get('session_q', '').strip()
    selected_session_id = request.GET.get('session', '').strip()
    query = request.GET.get('q', '').strip()
    result_filter = request.GET.get('result', '').strip()

    sessions = get_inventory_report_sessions_queryset(session_search)
    selected_session = sessions.filter(pk=selected_session_id).first() if selected_session_id else None
    if selected_session is None and selected_session_id:
        selected_session = InventorySession.objects.select_related('location', 'created_by').filter(pk=selected_session_id).first()

    items = get_inventory_report_items_queryset(selected_session, query, result_filter) if selected_session else InventoryItem.objects.none()
    stats = build_inventory_report_context(selected_session, items)

    query_params = request.GET.copy()
    export_query_string = query_params.urlencode()

    context = {
        'report_title': 'Отчёт по инвентаризации',
        'sessions': sessions,
        'selected_session': selected_session,
        'selected_session_id': selected_session_id,
        'selected_session_search': session_search,
        'items': items,
        'selected_query': query,
        'selected_result': result_filter,
        'result_choices': InventoryItem.ResultChoices.choices,
        'is_admin': is_admin,
        'is_manager': is_manager,
        'can_manage_reports': profile.can_manage_reports(),
        'export_query_string': export_query_string,
        **stats,
    }

    if selected_session:
        log_action(user, ActionLog.ActionType.REPORT_CREATE, f'Создан отчёт по инвентаризации: {selected_session.title}', selected_session.pk)

    return render(request, 'report_inventory_session.html', context)


def export_report_assets_by_employee_excel(request):
    """Экспорт отчёта 'Какие активы у сотрудников' в Excel"""
    
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error
    
    # Получаем данные
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(employee__isnull=False).order_by('employee__last_name', 'employee__first_name')
    
    # Создаём workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Активы сотрудников"
    
    # Стили
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Заголовки
    headers = ['Сотрудник', 'Наименование', 'Инв. номер', 'Тип', 'Локация', 'Статус']
    ws.append(headers)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    # Данные
    for asset in assets:
        employee_name = f"{asset.employee.last_name} {asset.employee.first_name}".strip() or asset.employee.username
        ws.append([
            employee_name,
            asset.asset_name,
            asset.inventory_number,
            asset.asset_type.name if asset.asset_type else '',
            asset.location.name if asset.location else '',
            asset.status.name if asset.status else '',
        ])
    
    # Применяем стили к данным
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
    
    # Автоширина
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 30)
    
    # Формируем ответ
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="report_assets_by_employee.xlsx"'
    wb.save(response)
    return response


def export_report_inventory_session_excel(request):
    """Экспорт отчёта по выбранной инвентаризации в Excel."""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error

    session_id = request.GET.get('session', '').strip()
    if not session_id:
        return HttpResponse('Не выбрана инвентаризация для экспорта', status=400)

    session = get_object_or_404(InventorySession.objects.select_related('location', 'created_by'), pk=session_id)
    query = request.GET.get('q', '').strip()
    result_filter = request.GET.get('result', '').strip()
    items = get_inventory_report_items_queryset(session, query, result_filter)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Инвентаризация'

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    meta_rows = [
        ['Отчёт', 'Инвентаризация'],
        ['Инвентаризация', session.title],
        ['Статус', session.get_status_display()],
        ['Локация', session.location.name if session.location else 'Все локации'],
        ['Создал', session.created_by.get_full_name() or session.created_by.username if session.created_by else '—'],
        ['Поиск', query or '—'],
        ['Фильтр результата', dict(InventoryItem.ResultChoices.choices).get(result_filter, 'Все')],
    ]
    for row in meta_rows:
        ws.append(row)
    ws.append([])

    headers = [
        'Инв. номер', 'Актив', 'Результат', 'Ожидался у сотрудника', 'Фактически у сотрудника',
        'Ожидаемая локация', 'Фактическая локация', 'Ожидаемый статус', 'Фактический статус',
        'Проверил', 'Дата проверки', 'Комментарий'
    ]
    ws.append(headers)
    header_row_index = ws.max_row

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row_index, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for item in items:
        ws.append([
            item.asset.inventory_number or '',
            item.asset.asset_name or '',
            item.get_result_display(),
            (item.expected_employee.get_full_name() or item.expected_employee.username) if item.expected_employee else '',
            (item.actual_employee.get_full_name() or item.actual_employee.username) if item.actual_employee else '',
            item.expected_location.name if item.expected_location else '',
            item.actual_location.name if item.actual_location else '',
            item.expected_status.name if item.expected_status else '',
            item.actual_status.name if item.actual_status else '',
            (item.checked_by.get_full_name() or item.checked_by.username) if item.checked_by else '',
            item.checked_at.strftime('%d.%m.%Y %H:%M') if item.checked_at else '',
            item.comment or '',
        ])

    for row in ws.iter_rows(min_row=header_row_index + 1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 28)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="inventory_report_{session.pk}.xlsx"'
    wb.save(response)
    return response


def export_report_assets_by_department_excel(request):
    """Экспорт отчёта 'Количество активов по отделам' в Excel"""
    
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error
    
    # Формируем данные
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(employee__isnull=False, location__isnull=False)
    
    report_data = {}
    for asset in assets:
        location_name = asset.location.name if asset.location else 'Без локации'
        type_name = asset.asset_type.name if asset.asset_type else 'Без типа'
        status_name = asset.status.name if asset.status else 'Без статуса'
        
        if location_name not in report_data:
            report_data[location_name] = {}
        if type_name not in report_data[location_name]:
            report_data[location_name][type_name] = {}
        if status_name not in report_data[location_name][type_name]:
            report_data[location_name][type_name][status_name] = 0
        report_data[location_name][type_name][status_name] += 1
    
    # Создаём workbook
    wb = Workbook()
    
    # Стили
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Создаём лист для каждого отдела
    for location_name, types in sorted(report_data.items()):
        # Ограничиваем имя листа
        safe_name = location_name[:30]
        ws = wb.create_sheet(title=safe_name)
        
        # Заголовки
        headers = ['Тип актива', 'В эксплуатации', 'На ремонте', 'Другие статусы', 'Итого']
        ws.append(headers)
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        # Данные
        for type_name, statuses in sorted(types.items()):
            in_use = statuses.get('В эксплуатации', 0)
            on_repair = statuses.get('На ремонте', 0)
            other = sum(count for status, count in statuses.items() if status not in ['В эксплуатации', 'На ремонте'])
            total = sum(statuses.values())
            
            ws.append([type_name, in_use, on_repair, other, total])
        
        # Применяем стили
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = border
        
        # Автоширина
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 20)
    
    # Удаляем default sheet если он пустой
    if wb.active.title == "Sheet" and wb.active.max_row == 1:
        del wb['Sheet']
    
    # Формируем ответ
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="report_assets_by_department.xlsx"'
    wb.save(response)
    return response


def export_report_assets_by_employee_pdf(request):
    """Экспорт отчёта 'Какие активы у сотрудников' в PDF"""
    
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error
    
    # Проверяем доступность ReportLab
    if not REPORTLAB_AVAILABLE:
        return HttpResponse("ReportLab не установлен", status=500)
    
    # Получаем данные
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(employee__isnull=False).order_by('employee__last_name', 'employee__first_name')
    
    # Создаём PDF
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="report_assets_by_employee.pdf"'
    
    # Шрифт с поддержкой кириллицы
    font_name = 'Arial' if font_registered else 'Helvetica'
    
    doc = SimpleDocTemplate(response, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    
    # Элементы PDF
    elements = []
    styles = getSampleStyleSheet()
    
    # Заголовок со шрифтом для кириллицы
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=16, spaceAfter=20, alignment=1, fontName=font_name)
    elements.append(Paragraph("Отчёт: Какие активы у сотрудников", title_style))
    
    # Таблица
    data = [['Сотрудник', 'Наименование', 'Инв. номер', 'Тип', 'Локация', 'Статус']]
    for asset in assets:
        employee_name = f"{asset.employee.last_name} {asset.employee.first_name}".strip() or asset.employee.username
        data.append([
            employee_name,
            asset.asset_name[:30] if asset.asset_name else '',
            asset.inventory_number or '',
            asset.asset_type.name if asset.asset_type else '',
            asset.location.name if asset.location else '',
            asset.status.name if asset.status else '',
        ])
    
    # Стили таблицы со шрифтом для кириллицы
    table = Table(data, colWidths=[40*mm, 35*mm, 25*mm, 25*mm, 25*mm, 25*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066CC')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), font_name),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F0F0')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)
    
    doc.build(elements)
    return response


def export_report_inventory_session_pdf(request):
    """Экспорт отчёта по выбранной инвентаризации в PDF."""
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error

    if not REPORTLAB_AVAILABLE:
        return HttpResponse("ReportLab не установлен", status=500)

    session_id = request.GET.get('session', '').strip()
    if not session_id:
        return HttpResponse('Не выбрана инвентаризация для экспорта', status=400)

    session = get_object_or_404(InventorySession.objects.select_related('location', 'created_by'), pk=session_id)
    query = request.GET.get('q', '').strip()
    result_filter = request.GET.get('result', '').strip()
    items = get_inventory_report_items_queryset(session, query, result_filter)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="inventory_report_{session.pk}.pdf"'

    font_name = 'Arial' if font_registered else 'Helvetica'
    doc = SimpleDocTemplate(response, pagesize=A4, leftMargin=10*mm, rightMargin=10*mm, topMargin=12*mm, bottomMargin=12*mm)

    elements = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=15, spaceAfter=10, alignment=1, fontName=font_name)
    meta_style = ParagraphStyle('Meta', parent=styles['BodyText'], fontSize=9, leading=12, fontName=font_name)

    elements.append(Paragraph(f"Отчёт по инвентаризации: {session.title}", title_style))
    elements.append(Paragraph(
        f"Статус: {session.get_status_display()}<br/>"
        f"Локация: {session.location.name if session.location else 'Все локации'}<br/>"
        f"Поиск: {query or '—'}<br/>"
        f"Фильтр результата: {dict(InventoryItem.ResultChoices.choices).get(result_filter, 'Все')}",
        meta_style
    ))
    elements.append(Spacer(1, 5*mm))

    data = [[
        'Инв. №', 'Актив', 'Результат', 'Ожидалось', 'Фактически', 'Комментарий'
    ]]
    for item in items:
        expected_employee = (item.expected_employee.get_full_name() or item.expected_employee.username) if item.expected_employee else '—'
        actual_employee = (item.actual_employee.get_full_name() or item.actual_employee.username) if item.actual_employee else '—'
        data.append([
            item.asset.inventory_number or '',
            (item.asset.asset_name or '')[:28],
            item.get_result_display(),
            expected_employee[:22],
            actual_employee[:22],
            (item.comment or '')[:40],
        ])

    table = Table(data, colWidths=[22*mm, 38*mm, 26*mm, 34*mm, 34*mm, 42*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066CC')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), font_name),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F0F0')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    return response


def export_report_assets_by_department_pdf(request):
    """Экспорт отчёта 'Количество активов по отделам' в PDF"""
    
    user = request.user
    profile = get_user_access(user)
    access_error = require_full_access(profile, 'reports', "Недостаточно прав для экспорта отчётов")
    if access_error:
        return access_error
    
    # Проверяем доступность ReportLab
    if not REPORTLAB_AVAILABLE:
        return HttpResponse("ReportLab не установлен", status=500)
    
    # Формируем данные
    assets = Asset.objects.select_related(
        'asset_type', 'status', 'employee', 'location'
    ).filter(employee__isnull=False, location__isnull=False)
    
    report_data = {}
    for asset in assets:
        location_name = asset.location.name if asset.location else 'Без локации'
        type_name = asset.asset_type.name if asset.asset_type else 'Без типа'
        status_name = asset.status.name if asset.status else 'Без статуса'
        
        if location_name not in report_data:
            report_data[location_name] = {}
        if type_name not in report_data[location_name]:
            report_data[location_name][type_name] = {}
        if status_name not in report_data[location_name][type_name]:
            report_data[location_name][type_name][status_name] = 0
        report_data[location_name][type_name][status_name] += 1
    
    # Создаём PDF
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="report_assets_by_department.pdf"'
    
    # Шрифт с поддержкой кириллицы
    font_name = 'Arial' if font_registered else 'Helvetica'
    
    doc = SimpleDocTemplate(response, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
    
    # Элементы PDF
    elements = []
    styles = getSampleStyleSheet()
    
    # Заголовок со шрифтом для кириллицы
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=16, spaceAfter=20, alignment=1, fontName=font_name)
    elements.append(Paragraph("Отчёт: Количество активов по отделам", title_style))
    
    # Стиль для заголовков отделов
    heading2_style = ParagraphStyle('Heading2', parent=styles['Heading2'], fontName=font_name)
    
    # Данные по отделам
    for location_name, types in sorted(report_data.items()):
        elements.append(Paragraph(f"<b>{location_name}</b>", heading2_style))
        elements.append(Spacer(1, 5*mm))
        
        data = [['Тип актива', 'В эксплуатации', 'На ремонте', 'Другие статусы', 'Итого']]
        for type_name, statuses in sorted(types.items()):
            in_use = statuses.get('В эксплуатации', 0)
            on_repair = statuses.get('На ремонте', 0)
            other = sum(count for status, count in statuses.items() if status not in ['В эксплуатации', 'На ремонте'])
            total = sum(statuses.values())
            
            data.append([type_name, in_use, on_repair, other, total])
        
        # Стили таблицы со шрифтом для кириллицы
        table = Table(data, colWidths=[50*mm, 25*mm, 25*mm, 25*mm, 25*mm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066CC')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), font_name),
            ('FONTNAME', (0, 1), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F0F0')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 10*mm))
    
    doc.build(elements)
    return response
